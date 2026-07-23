"""
Парсер для автоматического входа в личный кабинет
с использованием распознавания изображений для CAdES
"""

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import NoSuchElementException, TimeoutException
import time
import os
import sys
import logging
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import cv2
import numpy as np
import pyautogui
import mss
import ctypes

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)


def get_short_path(path):
    """Получение короткого пути (8.3 формат) для Windows"""
    try:
        buffer = ctypes.create_unicode_buffer(260)
        ctypes.windll.kernel32.GetShortPathNameW(path, buffer, 260)
        return buffer.value if buffer.value else path
    except:
        return path


def get_asset_path(filename):
    """Получение пути к файлу в папке assets"""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    # Используем короткий путь для обхода проблем с кириллицей
    short_script_dir = get_short_path(script_dir)
    assets_dir = os.path.join(short_script_dir, "assets")
    if not os.path.exists(assets_dir):
        os.makedirs(assets_dir)
    return os.path.join(assets_dir, filename)


class RoskaznaParser:
    URL = os.getenv("TARGET_URL")
    
    # Константные данные для Excel
    SKZI_TYPE = "КриптоПро CSP 5.0 ***** R3 *** ******-007671"
    SERVICE_TYPE = "Установка ключевых документов"
    KEY_TYPE = "ЭЦП"
    
    def __init__(self):
        self.yandex_path = r"C:\Program Files\Yandex\YandexBrowser\Application\browser.exe"
        print(f"🔍 Яндекс Браузер: {self.yandex_path}")
        self.driver = None
        self.all_records = []
        # Координаты кнопок (нужно подобрать под ваш экран)
        self.select_button_coords = None
        self.ok_button_coords = None
    
    def _find_chromedriver(self):
        possible_paths = [
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "chromedriver.exe"),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "chromedriver-win64", "chromedriver.exe"),
        ]
        
        for path in possible_paths:
            if os.path.exists(path):
                print(f"✅ Найден ChromeDriver: {path}")
                return path
        
        raise FileNotFoundError("chromedriver.exe не найден!")
    
    def _setup_driver(self):
        print("⚙️ Настройка WebDriver...")
        
        options = Options()
        options.binary_location = self.yandex_path
        
        options.add_argument('--disable-blink-features=AutomationControlled')
        options.add_experimental_option('excludeSwitches', ['enable-automation'])
        options.add_experimental_option('useAutomationExtension', False)
        options.add_argument('--disable-web-security')
        options.add_argument('--allow-running-insecure-content')
        options.add_argument('--disable-extensions')
        options.add_argument('--no-sandbox')
        options.add_argument('--enable-plugins')
        
        prefs = {
            "credentials_enable_service": True,
            "profile.password_manager_enabled": True,
            "client_certificate_portal": True,
        }
        options.add_experimental_option("prefs", prefs)
        
        chromedriver_path = self._find_chromedriver()
        service = Service(chromedriver_path)
        
        self.driver = webdriver.Chrome(service=service, options=options)
        self.driver.maximize_window()
        self.wait = WebDriverWait(self.driver, 10)
        print("✅ WebDriver готов")
    
    def open_site(self):
        print(f"🌐 Открываем: {self.URL}")
        self.driver.get(self.URL)
        time.sleep(3)
        return True
    
    def click_my_requests(self):
        print("📋 Поиск 'Мои запросы...'")
        
        js_script = """
        var items = document.querySelectorAll('span');
        for (var i = 0; i < items.length; i++) {
            if (items[i].textContent.includes('Мои запросы и отзыв сертификата')) {
                var parent = items[i].closest('div[class*="_menu-item"]');
                if (parent) {
                    parent.click();
                    return true;
                }
                items[i].click();
                return true;
            }
        }
        return false;
        """
        
        try:
            result = self.driver.execute_script(js_script)
            if result:
                print("✅ Кликнули по меню")
                time.sleep(2)
                return True
        except:
            pass
        
        print("❌ Не найден пункт меню")
        return False
    
    def click_certificate_login(self):
        print("🔐 Поиск 'Вход по сертификату'...")
        time.sleep(2)
        
        js_script = """
        var tiles = document.querySelectorAll('div[class*="_tile"]');
        for (var i = 0; i < tiles.length; i++) {
            if (tiles[i].textContent.includes('Вход по сертификату')) {
                tiles[i].click();
                return true;
            }
        }
        return false;
        """
        
        try:
            result = self.driver.execute_script(js_script)
            if result:
                print("✅ Кликнули по входу по сертификату")
                time.sleep(2)
                return True
        except:
            pass
        
        print("❌ Не найден 'Вход по сертификату'")
        return False
    
    def find_and_click_image(self, template_path, confidence=0.7, timeout=10, click_offset=(0, 0)):
        """
        Поиск изображения на экране и клик по нему
        """
        print(f"🔍 Ищем изображение: {os.path.basename(template_path)}")
        
        if not os.path.exists(template_path):
            print(f"❌ Файл не найден: {template_path}")
            return False
        
        template = cv2.imread(template_path, cv2.IMREAD_GRAYSCALE)
        if template is None:
            print(f"❌ Не удалось загрузить изображение: {template_path}")
            return False
        
        template_height, template_width = template.shape
        
        start_time = time.time()
        
        with mss.mss() as sct:
            monitor = sct.monitors[1]
            
            while time.time() - start_time < timeout:
                screenshot = sct.grab(monitor)
                screenshot_np = np.array(screenshot)
                screenshot_gray = cv2.cvtColor(screenshot_np, cv2.COLOR_BGR2GRAY)
                
                result = cv2.matchTemplate(screenshot_gray, template, cv2.TM_CCOEFF_NORMED)
                min_val, max_val, min_loc, max_loc = cv2.minMaxLoc(result)
                
                if max_val >= confidence:
                    center_x = max_loc[0] + template_width // 2 + click_offset[0]
                    center_y = max_loc[1] + template_height // 2 + click_offset[1]
                    
                    print(f"✅ Найдено! Уверенность: {max_val:.3f}, координаты: ({center_x}, {center_y})")
                    pyautogui.click(center_x, center_y)
                    print("✅ Клик выполнен")
                    return True
                
                time.sleep(0.3)
        
        print(f"❌ Изображение не найдено за {timeout} секунд")
        return False
    
    def get_coordinates_interactive(self, message):
        """Интерактивное получение координат от пользователя"""
        print(f"\n{message}")
        print("Наведите курсор на нужную кнопку и нажмите Enter...")
        input()
        x, y = pyautogui.position()
        print(f"✅ Координаты: ({x}, {y})")
        return x, y
    
    def click_select_certificate(self):
        """
        Нажатие кнопки 'Выбрать' через координаты или распознавание
        """
        print("🖱️ Нажимаем кнопку 'Выбрать'...")
        
        # Сначала пробуем найти кнопку через распознавание
        select_button_path = get_asset_path("select_button.png")
        if os.path.exists(select_button_path):
            if self.find_and_click_image(select_button_path, confidence=0.5, timeout=3):
                print("✅ Кликнули по кнопке 'Выбрать' через распознавание")
                time.sleep(1)
                return True
        
        # Если не нашли - используем координаты
        print("ℹ️ Используем координаты для клика по кнопке 'Выбрать'")
        
        # Координаты для кнопки "Выбрать" (подобраны для вашего разрешения)
        # Можно изменить, запустив get_coordinates_interactive()
        click_x = 960   # Центр экрана по X
        click_y = 640   # Центр экрана по Y (чуть ниже середины)
        
        print(f"🖱️ Кликаем по координатам: ({click_x}, {click_y})")
        pyautogui.click(click_x, click_y)
        print("✅ Клик выполнен")
        time.sleep(2)
        return True
    
    def handle_cades_window(self):
        """
        Обработка окна CAdES через распознавание изображений
        """
        print("⏳ Ожидание окна CAdES...")
        time.sleep(3)
        
        # Путь к изображению кнопки ОК
        ok_button_path = get_asset_path("ok_button.png")
        
        print(f"📸 Ищем кнопку 'ОК'...")
        
        # Ищем и кликаем по кнопке ОК через распознавание
        if os.path.exists(ok_button_path):
            if self.find_and_click_image(ok_button_path, confidence=0.5, timeout=8):
                print("✅ Кнопка 'ОК' нажата через распознавание!")
                time.sleep(2)
                return True
        
        # Если не нашли - пробуем нажать Enter
        print("🔄 Пробуем нажать Enter...")
        try:
            import keyboard
            keyboard.press_and_release('enter')
            print("✅ Enter нажат")
            time.sleep(2)
            return True
        except:
            pass
        
        # Если не работает - пробуем pyautogui Enter
        try:
            pyautogui.press('enter')
            print("✅ Enter нажат через pyautogui")
            time.sleep(2)
            return True
        except:
            pass
        
        # Если ничего не работает
        print("⚠️ Не удалось найти кнопку 'ОК' автоматически")
        print("👉 Пожалуйста, нажмите 'ОК' вручную в окне CAdES")
        input("Нажмите Enter после нажатия 'ОК'...")
        
        return True
    
    def wait_for_login(self, timeout=10):
        print("⏳ Ожидание входа в личный кабинет...")
        time.sleep(timeout)
        return True
    
    def _parse_date(self, date_str):
        try:
            date_str = date_str.strip()
            return datetime.strptime(date_str, "%d.%m.%Y, %H:%M")
        except:
            try:
                return datetime.strptime(date_str, "%d.%m.%Y")
            except:
                return None
    
    def scroll_and_collect(self, max_records=999999, days_back=15):
        print(f"\n📊 Начинаем сбор данных за последние {days_back} дней...")
        
        cutoff_date = datetime.now() - timedelta(days=days_back)
        print(f"📅 Собираем записи с {cutoff_date.strftime('%d.%m.%Y')}")
        
        collected = 0
        previous_count = 0
        scroll_attempts = 0
        max_scroll_attempts = 100
        no_new_records_count = 0
        max_no_new = 10
        stop_collecting = False
        
        time.sleep(2)
        
        while scroll_attempts < max_scroll_attempts and not stop_collecting:
            try:
                items = self.driver.find_elements(By.CSS_SELECTOR, "div._list-item_ycb2i_14")
                current_count = len(items)
                
                if current_count > previous_count:
                    new_count = current_count - previous_count
                    print(f"\n📥 Найдено {new_count} новых записей (всего: {current_count})")
                    
                    for i in range(previous_count, current_count):
                        try:
                            record = self._extract_record_from_item(items[i])
                            if record and record.get('serial_number') and record.get('date_from'):
                                record_date = self._parse_date(record.get('date_from', ''))
                                
                                if record_date:
                                    if record_date < cutoff_date:
                                        print(f"\n⏹️ Достигнута дата {record_date.strftime('%d.%m.%Y')} (раньше {cutoff_date.strftime('%d.%m.%Y')})")
                                        print("📌 Останавливаем сбор данных")
                                        stop_collecting = True
                                        break
                                    
                                    self.all_records.append(record)
                                    collected += 1
                                    
                                    if collected % 5 == 0:
                                        print(f"  ✅ Запись #{collected}: {record.get('fio', 'Без ФИО')[:30]}... - {record.get('date_from', '')}")
                                else:
                                    print(f"  ⚠️ Не удалось распарсить дату: {record.get('date_from', '')}")
                                    
                        except Exception as e:
                            print(f"  ⚠️ Ошибка при обработке записи {i+1}: {e}")
                    
                    previous_count = current_count
                    no_new_records_count = 0
                else:
                    no_new_records_count += 1
                    print(f"⏳ Ожидание загрузки... ({no_new_records_count}/{max_no_new})")
                    
                    if no_new_records_count >= max_no_new:
                        print("\n✅ Все записи загружены!")
                        break
                
                if stop_collecting:
                    break
                
                print(f"📜 Прокрутка страницы вниз (попытка {scroll_attempts + 1})...")
                
                try:
                    if items:
                        last_item = items[-1]
                        self.driver.execute_script("arguments[0].scrollIntoView(true);", last_item)
                        self.driver.execute_script("window.scrollBy(0, 500);")
                    else:
                        self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                    self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                except:
                    self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                
                time.sleep(2.5)
                scroll_attempts += 1
                
                if len(self.all_records) > 0 and len(self.all_records) % 20 == 0:
                    print(f"📊 Прогресс: собрано {len(self.all_records)} записей")
                    
            except Exception as e:
                print(f"⚠️ Ошибка при сборе данных: {e}")
                time.sleep(1)
                scroll_attempts += 1
        
        print(f"\n📊 Итого собрано записей: {len(self.all_records)} за последние {days_back} дней")
        return self.all_records
    
    def _extract_record_from_item(self, item):
        record = {}
        
        try:
            try:
                date_from_elem = item.find_element(By.CSS_SELECTOR, "p._date__from_ycb2i_47")
                record['date_from'] = date_from_elem.text.replace("от ", "").strip()
            except:
                record['date_from'] = ""
            
            try:
                all_p = item.find_elements(By.TAG_NAME, "p")
                for p in all_p:
                    text = p.text.strip()
                    if (text and 
                        not text.startswith("от ") and 
                        not text.startswith("до ") and
                        not text.startswith("СНИЛС:") and
                        len(text) > 5 and
                        not text.isdigit() and
                        len(text) < 50):
                        record['fio'] = text
                        break
            except:
                record['fio'] = ""
            
            try:
                info_divs = item.find_elements(By.CSS_SELECTOR, "div._info_ycb2i_37")
                for div in info_divs:
                    try:
                        p_text = div.find_element(By.TAG_NAME, "p").text.strip()
                        if not p_text.startswith("СНИЛС:") and len(p_text) > 10:
                            record['serial_number'] = p_text
                            break
                    except:
                        continue
            except:
                record['serial_number'] = ""
            
            try:
                date_to_elem = item.find_element(By.XPATH, ".//p[contains(text(), 'до ')]")
                record['date_to'] = date_to_elem.text.replace("до ", "").strip()
            except:
                record['date_to'] = ""
            
            return record
            
        except Exception as e:
            return None
    
    def create_excel(self, filename=None):
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"СКЗИ_отчет_{timestamp}.xlsx"
        
        print(f"\n📝 Создание Excel файла: {filename}")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "СКЗИ"
        
        times_font = Font(name='Times New Roman', size=11)
        bold_times = Font(name='Times New Roman', size=11, bold=True)
        bold_number_font = Font(name='Times New Roman', size=11, bold=True)
        
        headers = [
            "№ п/п",
            "Дата установки",
            "Тип и серийные номера используемых СКЗИ",
            "Записи по обслуживанию СКЗИ",
            "ФИО пользователя",
            "Срок действия",
            "Используемые криптоключи\nТип ключевого документа",
            "Используемые криптоключи\nСерийный, криптографический номер и номер экземпляра ключевого документа",
            "Используемые криптоключи\nНомер разового ключевого носителя или зоны СКЗИ, в которую введены криптоключи",
            "Отметка об уничтожении(стирании)\nДата",
            "Отметка об уничтожении(стирании)\nПодпись пользователя СКЗИ",
            "Примечание"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = bold_times
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
        
        column_widths = {
            1: 8, 2: 18, 3: 45, 4: 30, 5: 30, 6: 20,
            7: 25, 8: 45, 9: 35, 10: 15, 11: 25, 12: 20
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        
        print(f"📊 Запись данных в Excel ({len(self.all_records)} записей)...")
        
        for idx, record in enumerate(self.all_records, start=1):
            row = idx + 1
            
            cells_data = [
                idx,
                record.get("date_from", ""),
                self.SKZI_TYPE,
                self.SERVICE_TYPE,
                record.get("fio", ""),
                record.get("date_to", ""),
                self.KEY_TYPE,
                record.get("serial_number", ""),
                "",
                "",
                "",
                ""
            ]
            
            for col, value in enumerate(cells_data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                if col == 1:
                    cell.font = bold_number_font
                else:
                    cell.font = times_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            if idx % 10 == 0 or idx == len(self.all_records):
                print(f"  Записано {idx}/{len(self.all_records)} записей...")
        
        ws.row_dimensions[1].height = 60
        for row in range(2, len(self.all_records) + 2):
            ws.row_dimensions[row].height = 25
        
        wb.save(filename)
        print(f"✅ Excel файл сохранен: {filename}")
        os.startfile(os.path.dirname(os.path.abspath(filename)))
        
        return filename
    
    def run(self, max_records=999999, days_back=15):
        print("\n" + "="*60)
        print("🚀 ЗАПУСК ПАРСЕРА")
        print("="*60 + "\n")
        
        try:
            self._setup_driver()
            self.open_site()
            self.click_my_requests()
            self.click_certificate_login()
            
            # Нажимаем кнопку "Выбрать"
            self.click_select_certificate()
            
            # Обрабатываем окно CAdES
            self.handle_cades_window()
            
            print("⏳ Ожидание входа в личный кабинет...")
            time.sleep(8)
            
            self.scroll_and_collect(max_records=max_records, days_back=days_back)
            
            if self.all_records:
                print(f"\n📊 Всего собрано записей: {len(self.all_records)}")
                excel_file = self.create_excel()
                print(f"✅ Данные сохранены в: {excel_file}")
            else:
                print("⚠️ Нет данных для сохранения")
            
            print("\n" + "="*60)
            print("✅ ПАРСИНГ ЗАВЕРШЕН!")
            print("="*60)
            input("\nНажмите Enter для завершения...")
            return True
            
        except Exception as e:
            print(f"\n❌ Ошибка: {e}")
            import traceback
            traceback.print_exc()
            return False
        finally:
            if self.driver:
                self.driver.quit()


if __name__ == "__main__":
    parser = RoskaznaParser()
    parser.run(max_records=999999, days_back=15)