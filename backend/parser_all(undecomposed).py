"""
Парсер для автоматического входа в личный кабинет
и сохранения данных в Excel
"""

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import NoSuchElementException, TimeoutException
import time
import os
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)


class ProgressBar:
    """Класс для отображения прогресса в консоли"""
    
    def __init__(self, total, prefix='', suffix='', decimals=1, length=50, fill='█', print_end='\r'):
        self.total = total
        self.prefix = prefix
        self.suffix = suffix
        self.decimals = decimals
        self.length = length
        self.fill = fill
        self.print_end = print_end
        self.iteration = 0
    
    def update(self, iteration=None):
        if iteration is not None:
            self.iteration = iteration
        else:
            self.iteration += 1
        
        percent = ("{0:." + str(self.decimals) + "f}").format(100 * (self.iteration / float(self.total)))
        filled_length = int(self.length * self.iteration // self.total)
        bar = self.fill * filled_length + '-' * (self.length - filled_length)
        
        print(f'\r{self.prefix} |{bar}| {percent}% {self.suffix}', end=self.print_end)
        
        if self.iteration == self.total:
            print()
    
    def complete(self, message='✅ Готово!'):
        self.update(self.total)
        print(f'  {message}')


class RoskaznaParser:
    URL = os.getenv("TARGET_URL")
    
    # Константные данные для Excel
    SKZI_TYPE = "КриптоПро CSP 5.0 ***** R3 *** ******-007671"
    SERVICE_TYPE = "Установка ключевых документов"
    KEY_TYPE = "ЭЦП"
    
    def __init__(self):
        self.yandex_path = r"C:\Program Files\Yandex\YandexBrowser\Application\browser.exe"
        print(f"🔍 Яндекс Браузер: {self.yandex_path}")
        self.driver = None
        self.all_records = []
    
    def _find_chromedriver(self):
        possible_paths = [
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "chromedriver.exe"),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "chromedriver-win64", "chromedriver.exe"),
        ]
        
        for path in possible_paths:
            if os.path.exists(path):
                print(f"✅ Найден ChromeDriver: {path}")
                return path
        
        raise FileNotFoundError("chromedriver.exe не найден!")
    
    def _setup_driver(self):
        print("⚙️ Настройка WebDriver...")
        
        options = Options()
        options.binary_location = self.yandex_path
        
        options.add_argument('--disable-blink-features=AutomationControlled')
        options.add_experimental_option('excludeSwitches', ['enable-automation'])
        options.add_experimental_option('useAutomationExtension', False)
        options.add_argument('--disable-web-security')
        options.add_argument('--allow-running-insecure-content')
        options.add_argument('--disable-extensions')
        options.add_argument('--no-sandbox')
        options.add_argument('--enable-plugins')
        
        prefs = {
            "credentials_enable_service": True,
            "profile.password_manager_enabled": True,
            "client_certificate_portal": True,
        }
        options.add_experimental_option("prefs", prefs)
        
        chromedriver_path = self._find_chromedriver()
        service = Service(chromedriver_path)
        
        self.driver = webdriver.Chrome(service=service, options=options)
        self.driver.maximize_window()
        self.wait = WebDriverWait(self.driver, 10)
        print("✅ WebDriver готов")
    
    def open_site(self):
        print(f"🌐 Открываем: {self.URL}")
        self.driver.get(self.URL)
        time.sleep(3)
        return True
    
    def click_my_requests(self):
        print("📋 Поиск 'Мои запросы...'")
        
        js_script = """
        var items = document.querySelectorAll('span');
        for (var i = 0; i < items.length; i++) {
            if (items[i].textContent.includes('Мои запросы и отзыв сертификата')) {
                var parent = items[i].closest('div[class*="_menu-item"]');
                if (parent) {
                    parent.click();
                    return true;
                }
                items[i].click();
                return true;
            }
        }
        return false;
        """
        
        try:
            result = self.driver.execute_script(js_script)
            if result:
                print("✅ Кликнули по меню")
                time.sleep(2)
                return True
        except:
            pass
        
        print("❌ Не найден пункт меню")
        return False
    
    def click_certificate_login(self):
        print("🔐 Поиск 'Вход по сертификату'...")
        time.sleep(2)
        
        js_script = """
        var tiles = document.querySelectorAll('div[class*="_tile"]');
        for (var i = 0; i < tiles.length; i++) {
            if (tiles[i].textContent.includes('Вход по сертификату')) {
                tiles[i].click();
                return true;
            }
        }
        return false;
        """
        
        try:
            result = self.driver.execute_script(js_script)
            if result:
                print("✅ Кликнули по входу по сертификату")
                time.sleep(2)
                return True
        except:
            pass
        
        print("❌ Не найден 'Вход по сертификату'")
        return False
    
    def click_select_certificate(self):
        print("🖱️ Нажимаем кнопку 'Выбрать'...")
        
        js_script = """
        var buttons = document.querySelectorAll('button');
        for (var i = 0; i < buttons.length; i++) {
            if (buttons[i].textContent.includes('Выбрать')) {
                buttons[i].click();
                return true;
            }
        }
        return false;
        """
        
        try:
            result = self.driver.execute_script(js_script)
            if result:
                print("✅ Кликнули по кнопке 'Выбрать'")
                return True
        except:
            pass
        
        print("ℹ️ Кнопка не найдена, возможно уже нажата")
        return True
    
    def press_enter(self):
        try:
            import keyboard
            keyboard.press_and_release('enter')
            print("✅ Enter нажат через keyboard")
            return True
        except ImportError:
            pass
        
        try:
            import pyautogui
            pyautogui.press('enter')
            print("✅ Enter нажат через pyautogui")
            return True
        except ImportError:
            pass
        
        try:
            import ctypes
            ctypes.windll.user32.keybd_event(0x0D, 0, 0, 0)
            time.sleep(0.1)
            ctypes.windll.user32.keybd_event(0x0D, 0, 0x0002, 0)
            print("✅ Enter нажат через ctypes")
            return True
        except:
            pass
        
        print("⚠️ Не удалось нажать Enter автоматически")
        return False
    
    def wait_for_login(self, timeout=10):
        print("⏳ Ожидание входа в личный кабинет...")
        time.sleep(timeout)
        return True
    
    def scroll_and_collect(self, max_records=999999):
        """
        Прокрутка страницы и сбор данных с бесконечной загрузкой
        """
        print(f"\n📊 Начинаем сбор данных...")
        
        collected = 0
        previous_count = 0
        scroll_attempts = 0
        max_scroll_attempts = 100  # Увеличили до 100
        no_new_records_count = 0
        max_no_new = 10  # Увеличили до 10 попыток
        
        # Ждем загрузки первой партии
        time.sleep(2)
        
        while scroll_attempts < max_scroll_attempts:
            # Находим все записи
            items = self.driver.find_elements(By.CSS_SELECTOR, "div._list-item_ycb2i_14")
            current_count = len(items)
            
            # Если появились новые записи
            if current_count > previous_count:
                # Обрабатываем новые записи
                new_count = current_count - previous_count
                print(f"\n📥 Найдено {new_count} новых записей (всего: {current_count})")
                
                for i in range(previous_count, current_count):
                    try:
                        record = self._extract_record_from_item(items[i])
                        if record and record.get('serial_number'):
                            self.all_records.append(record)
                            collected += 1
                            # Показываем прогресс каждые 10 записей
                            if collected % 10 == 0:
                                print(f"  ✅ Запись #{collected}: {record.get('fio', 'Без ФИО')[:30]}... - {record.get('serial_number', 'Нет ключа')[:20]}...")
                    except Exception as e:
                        print(f"  ⚠️ Ошибка при обработке записи {i+1}: {e}")
                
                previous_count = current_count
                no_new_records_count = 0
            else:
                no_new_records_count += 1
                print(f"⏳ Ожидание загрузки... ({no_new_records_count}/{max_no_new})")
                
                if no_new_records_count >= max_no_new:
                    print("\n✅ Все записи загружены!")
                    break
            
            # Прокручиваем страницу вниз (более агрессивно)
            print(f"📜 Прокрутка страницы вниз (попытка {scroll_attempts + 1})...")
            
            try:
                # Прокручиваем к последней записи
                if items:
                    last_item = items[-1]
                    self.driver.execute_script("arguments[0].scrollIntoView(true);", last_item)
                    # Прокручиваем еще ниже
                    self.driver.execute_script("window.scrollBy(0, 500);")
                else:
                    self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                
                # Дополнительная прокрутка вниз
                self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                
            except:
                self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            
            time.sleep(2.5)  # Увеличили задержку для загрузки
            scroll_attempts += 1
            
            # Показываем общий прогресс каждые 50 записей
            if len(self.all_records) > 0 and len(self.all_records) % 50 == 0:
                print(f"📊 Прогресс: собрано {len(self.all_records)} записей")
        
        print(f"\n📊 Итого собрано записей: {len(self.all_records)}")
        return self.all_records
    
    def _extract_record_from_item(self, item):
        """
        Извлечение данных из элемента записи _list-item_ycb2i_14
        """
        record = {}
        
        try:
            # 1. Дата установки
            try:
                date_from_elem = item.find_element(By.CSS_SELECTOR, "p._date__from_ycb2i_47")
                record['date_from'] = date_from_elem.text.replace("от ", "").strip()
            except:
                record['date_from'] = ""
            
            # 2. ФИО пользователя
            try:
                all_p = item.find_elements(By.TAG_NAME, "p")
                for p in all_p:
                    text = p.text.strip()
                    if (text and 
                        not text.startswith("от ") and 
                        not text.startswith("до ") and
                        not text.startswith("СНИЛС:") and
                        len(text) > 5 and
                        not text.isdigit() and
                        len(text) < 50):
                        record['fio'] = text
                        break
            except:
                record['fio'] = ""
            
            # 3. Серийный номер ключа
            try:
                info_divs = item.find_elements(By.CSS_SELECTOR, "div._info_ycb2i_37")
                
                for div in info_divs:
                    try:
                        p_text = div.find_element(By.TAG_NAME, "p").text.strip()
                        if not p_text.startswith("СНИЛС:") and len(p_text) > 10:
                            record['serial_number'] = p_text
                            break
                    except:
                        continue
            except:
                record['serial_number'] = ""
            
            # 4. Срок действия
            try:
                date_to_elem = item.find_element(By.XPATH, ".//p[contains(text(), 'до ')]")
                record['date_to'] = date_to_elem.text.replace("до ", "").strip()
            except:
                record['date_to'] = ""
            
            return record
            
        except Exception as e:
            return None
    
    def create_excel(self, filename=None):
        """Создание Excel файла с данными в Times New Roman"""
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"СКЗИ_отчет_{timestamp}.xlsx"
        
        print(f"\n📝 Создание Excel файла: {filename}")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "СКЗИ"
        
        # Шрифт Times New Roman для всего документа
        times_font = Font(name='Times New Roman', size=11)
        bold_times = Font(name='Times New Roman', size=11, bold=True)
        
        # Заголовки
        headers = [
            "№ п/п",
            "Дата установки",
            "Тип и серийные номера используемых СКЗИ",
            "Записи по обслуживанию СКЗИ",
            "ФИО пользователя",
            "Срок действия",
            "Используемые криптоключи\nТип ключевого документа",
            "Используемые криптоключи\nСерийный, криптографический номер и номер экземпляра ключевого документа",
            "Используемые криптоключи\nНомер разового ключевого носителя или зоны СКЗИ, в которую введены криптоключи",
            "Отметка об уничтожении(стирании)\nДата",
            "Отметка об уничтожении(стирании)\nПодпись пользователя СКЗИ",
            "Примечание"
        ]
        
        # Записываем заголовки
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = bold_times
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            # Светло-серый фон для заголовков
            cell.fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
        
        # Устанавливаем ширину колонок
        column_widths = {
            1: 8, 2: 18, 3: 45, 4: 30, 5: 30, 6: 20,
            7: 25, 8: 45, 9: 35, 10: 15, 11: 25, 12: 20
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # Записываем данные
        print(f"📊 Запись данных в Excel ({len(self.all_records)} записей)...")
        
        for idx, record in enumerate(self.all_records, start=1):
            row = idx + 1
            
            # Данные
            cells_data = [
                idx,
                record.get("date_from", ""),
                self.SKZI_TYPE,
                self.SERVICE_TYPE,
                record.get("fio", ""),
                record.get("date_to", ""),
                self.KEY_TYPE,
                record.get("serial_number", ""),
                "",
                "",
                "",
                ""
            ]
            
            for col, value in enumerate(cells_data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = times_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            # Прогресс записи
            if idx % 10 == 0 or idx == len(self.all_records):
                print(f"  Записано {idx}/{len(self.all_records)} записей...")
        
        # Высота строк
        ws.row_dimensions[1].height = 40
        for row in range(2, len(self.all_records) + 2):
            ws.row_dimensions[row].height = 25
        
        # Сохраняем
        wb.save(filename)
        print(f"✅ Excel файл сохранен: {filename}")
        
        # Открываем папку
        os.startfile(os.path.dirname(os.path.abspath(filename)))
        
        return filename
    
    def run(self, max_records=999999):
        """Основной метод"""
        print("\n" + "="*60)
        print("🚀 ЗАПУСК ПАРСЕРА")
        print("="*60 + "\n")
        
        try:
            self._setup_driver()
            self.open_site()
            self.click_my_requests()
            self.click_certificate_login()
            self.click_select_certificate()
            
            print("\n⏳ Ожидание открытия CAdES...")
            time.sleep(2)
            
            print("⌨️ Нажимаем Enter...")
            self.press_enter()
            
            print("⏳ Ожидание входа в личный кабинет...")
            time.sleep(5)
            
            # Собираем все записи (без ограничения)
            self.scroll_and_collect(max_records=max_records)
            
            # Создаем Excel
            if self.all_records:
                print(f"\n📊 Всего собрано записей: {len(self.all_records)}")
                excel_file = self.create_excel()
                print(f"✅ Данные сохранены в: {excel_file}")
            else:
                print("⚠️ Нет данных для сохранения")
            
            print("\n" + "="*60)
            print("✅ ПАРСИНГ ЗАВЕРШЕН!")
            print("="*60)
            input("\nНажмите Enter для завершения...")
            return True
            
        except Exception as e:
            print(f"\n❌ Ошибка: {e}")
            import traceback
            traceback.print_exc()
            return False
        finally:
            if self.driver:
                self.driver.quit()


if __name__ == "__main__":
    parser = RoskaznaParser()
    # max_records=999999 - собираем все записи
    parser.run(max_records=999999)