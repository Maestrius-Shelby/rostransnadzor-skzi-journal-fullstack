"""
Экспорт данных в Excel и JSON
"""
import os
import json
import tempfile
import shutil
from datetime import datetime
from typing import List, Dict, Optional
import csv

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..utils.logger import get_logger
from ..utils.path_utils import get_output_path
from ..core.config import ExportConfig

logger = get_logger(__name__)


def is_file_locked(filepath: str) -> bool:
    """Проверить, заблокирован ли файл другим процессом"""
    if not os.path.exists(filepath):
        return False
    
    # Проверяем наличие lock-файлов LibreOffice
    lock_file = os.path.join(os.path.dirname(filepath), f".~lock.{os.path.basename(filepath)}#")
    if os.path.exists(lock_file):
        return True
    
    # Проверяем lock-файлы для других редакторов
    lock_files = [
        f"{filepath}.lock",
        f"{filepath}.lck",
        f".~lock.{os.path.basename(filepath)}",
        f".~lock.{os.path.basename(filepath)}#",  # LibreOffice
    ]
    
    for lock in lock_files:
        if os.path.exists(lock):
            return True
    
    # Windows: попытка открыть файл
    if platform.system() == 'Windows':
        try:
            with open(filepath, 'r+b') as f:
                pass
            return False
        except (IOError, OSError, PermissionError):
            return True
    
    # Linux / macOS: пробуем fcntl
    try:
        import fcntl
        with open(filepath, 'r+b') as f:
            fcntl.flock(f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            fcntl.flock(f.fileno(), fcntl.LOCK_UN)
            return False
    except (IOError, OSError, PermissionError, BlockingIOError):
        return True
    except ImportError:
        try:
            with open(filepath, 'r+b') as f:
                pass
            return False
        except (IOError, OSError, PermissionError):
            return True


def get_downloads_path(filename: str) -> str:
    """Получить путь к папке Downloads пользователя"""
    if os.name == 'nt':
        downloads = os.path.join(os.environ['USERPROFILE'], 'Downloads')
    else:
        downloads = os.path.join(os.path.expanduser('~'), 'Downloads')
    
    os.makedirs(downloads, exist_ok=True)
    return os.path.join(downloads, filename)


class ExcelExporter:
    """Экспорт данных"""
    
    def __init__(
        self,
        config: ExportConfig,
        skzi_type: str,
        service_type: str,
        key_type: str
    ):
        self.config = config
        self.skzi_type = skzi_type
        self.service_type = service_type
        self.key_type = key_type
        
        self.COLOR_NEW = "C6EFCE"
        self.COLOR_EXPIRING = "FFEB9C"
        self.COLOR_EXPIRED = "FFC7CE"
        self.COLOR_HEADER = "E8E8E8"
    
    def export_to_excel(
        self, 
        records: List[Dict[str, str]], 
        output_path: Optional[str] = None,
        highlight_new: bool = True,
        highlight_expiring: bool = True,
        highlight_expired: bool = True
    ) -> Dict[str, Optional[str]]:
        """
        Экспорт записей в Excel.
        
        - Всегда сохраняет в output_path (если не заблокирован)
        - Возвращает путь к файлу для скачивания
        """
        result = {'output': None, 'download': None}
        
        # Проверяем, не открыт ли файл
        if output_path and os.path.exists(output_path) and is_file_locked(output_path):
            print(f"⛔ Файл {os.path.basename(output_path)} ОТКРЫТ в Excel!")
            print(f"   💡 ЗАКРОЙТЕ файл и повторите экспорт")
            return result
        
        # Создаём workbook из records
        sorted_records = self._sort_records(records)
        wb = self._create_workbook(sorted_records, highlight_new, highlight_expiring, highlight_expired)
        
        # Сохраняем в output
        if output_path:
            output_dir = os.path.dirname(output_path)
            os.makedirs(output_dir, exist_ok=True)
            
            with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                temp_path = tmp.name
            try:
                wb.save(temp_path)
                shutil.copy2(temp_path, output_path)
                result['output'] = output_path
                result['download'] = output_path  # ✅ Для скачивания используем тот же файл
                print(f"📊 Файл сохранён: {output_path}")
            finally:
                if os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except:
                        pass
        
        # Статистика
        if result.get('output'):
            new_count = sum(1 for r in sorted_records if r.get('is_new', False))
            expiring_count = sum(1 for r in sorted_records if r.get('is_expiring', False))
            expired_count = sum(1 for r in sorted_records if r.get('is_expired', False))
            
            print(f"\n📊 Статистика экспорта:")
            print(f"   📝 Всего записей: {len(sorted_records)}")
            print(f"   🟢 Новых: {new_count}")
            print(f"   🟠 Истекающих: {expiring_count}")
            print(f"   🔴 Истекших: {expired_count}")
        
        return result

    def _create_workbook(self, records, highlight_new, highlight_expiring, highlight_expired):
        """Создать workbook с данными"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "СКЗИ"
        
        # Шрифты
        times_font = Font(name=self.config.excel_font, size=self.config.font_size)
        bold_times = Font(name=self.config.excel_font, size=self.config.font_size, bold=True)
        bold_number_font = Font(name=self.config.excel_font, size=self.config.font_size, bold=True)
        
        # Границы
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Заголовки
        headers = [
            "№ п/п", "Дата установки", "Тип и серийные номера используемых СКЗИ",
            "Записи по обслуживанию СКЗИ", "ФИО пользователя", "Срок действия",
            "Используемые криптоключи\nТип ключевого документа",
            "Используемые криптоключи\nСерийный, криптографический номер и номер экземпляра ключевого документа",
            "Используемые криптоключи\nНомер разового ключевого носителя или зоны СКЗИ, в которую введены криптоключи",
            "Отметка об уничтожении(стирании)\nДата",
            "Отметка об уничтожении(стирании)\nПодпись пользователя СКЗИ",
            "Примечание"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = bold_times
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.fill = PatternFill(start_color=self.COLOR_HEADER, end_color=self.COLOR_HEADER, fill_type="solid")
            cell.border = thin_border
        
        column_widths = {
            1: 8, 2: 18, 3: 45, 4: 30, 5: 30, 6: 20,
            7: 25, 8: 45, 9: 35, 10: 15, 11: 25, 12: 20
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # Заполняем данные
        for idx, record in enumerate(records, start=1):
            row = idx + 1
            
            is_new = record.get('is_new', False)
            is_expiring = record.get('is_expiring', False)
            is_expired = record.get('is_expired', False)
            
            cells_data = [
                idx,
                record.get("date_from", ""),
                self.skzi_type,
                self.service_type,
                record.get("fio", ""),
                record.get("date_to", ""),
                record.get("key_type", self.key_type),
                record.get("serial_number", ""),
                record.get("key_carrier_number", ""),
                record.get("destruction_date", ""),
                record.get("destruction_signature", ""),
                record.get("notes", "")
            ]
            
            fill_color = None
            if highlight_expired and is_expired:
                fill_color = self.COLOR_EXPIRED
            elif highlight_expiring and is_expiring:
                fill_color = self.COLOR_EXPIRING
            elif highlight_new and is_new:
                fill_color = self.COLOR_NEW
            
            for col, value in enumerate(cells_data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = bold_number_font if col == 1 else times_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border
                if fill_color:
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        
        ws.row_dimensions[1].height = self.config.header_height
        for row in range(2, len(records) + 2):
            ws.row_dimensions[row].height = self.config.row_height
        
        self._add_stats_sheet(wb, len(records), 
            sum(1 for r in records if r.get('is_new', False)),
            sum(1 for r in records if r.get('is_expiring', False)),
            sum(1 for r in records if r.get('is_expired', False))
        )
        
        return wb
    
    def _sort_records(self, records: List[Dict]) -> List[Dict]:
        """Сортировка записей"""
        def sort_key(record):
            is_expiring = record.get('is_expiring', False)
            is_expired = record.get('is_expired', False)
            is_new = record.get('is_new', False)
            
            if is_expiring:
                return 0
            elif is_expired:
                return 1
            elif is_new:
                return 2
            else:
                return 3
        
        return sorted(records, key=sort_key)
    
    def _add_stats_sheet(self, wb, total: int, new: int, expiring: int, expired: int):
        """Добавить лист со статистикой"""
        ws = wb.create_sheet("Статистика")
        
        ws['A1'] = "СТАТИСТИКА ПАРСИНГА"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:B1')
        
        ws['A2'] = f"Дата экспорта: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws['A2'].font = Font(size=11)
        ws.merge_cells('A2:B2')
        
        stats = [
            ("Всего записей", total, None),
            ("🟢 Новые записи", new, self.COLOR_NEW),
            ("🟠 Истекают (до 15 дней)", expiring, self.COLOR_EXPIRING),
            ("🔴 Истекли", expired, self.COLOR_EXPIRED)
        ]
        
        for i, (label, value, color) in enumerate(stats, start=4):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value
            ws[f'A{i}'].font = Font(size=12)
            ws[f'B{i}'].font = Font(size=12, bold=True)
            
            if color and value > 0:
                ws[f'A{i}'].fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                ws[f'B{i}'].fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 20
        
        # Легенда
        legend_start = 9
        ws[f'A{legend_start}'] = "ЛЕГЕНДА:"
        ws[f'A{legend_start}'].font = Font(size=12, bold=True)
        
        legend = [
            ("🟢 Зеленый", "Новые записи (добавлены при последнем парсинге)", self.COLOR_NEW),
            ("🟠 Оранжевый", "Истекают в ближайшие 15 дней", self.COLOR_EXPIRING),
            ("🔴 Красный", "Срок действия истек (запись удалена с сайта)", self.COLOR_EXPIRED)
        ]
        
        for i, (icon, text, color) in enumerate(legend, start=legend_start + 1):
            ws[f'A{i}'] = icon
            ws[f'B{i}'] = text
            ws[f'A{i}'].fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            ws[f'A{i}'].font = Font(size=11)
            ws[f'B{i}'].font = Font(size=11)
        
        ws.column_dimensions['B'].width = 55
    
    def export_to_json(self, records: List[Dict], filepath: str) -> str:
        """Экспорт в JSON"""
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump({
                'records': records,
                'last_updated': datetime.now().isoformat(),
                'total_count': len(records)
            }, f, ensure_ascii=False, indent=2)
        return filepath